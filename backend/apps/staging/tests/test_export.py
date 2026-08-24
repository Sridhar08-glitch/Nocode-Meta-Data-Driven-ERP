"""Export pipeline (PROJECT_HANDBOOK.md §29.1 / §29.4)."""
import io

import pytest

from apps.records.services import RecordService
from apps.staging.models import ExportJob
from apps.staging.services import ExportService, ImportError
from apps.staging.tasks import expire_export_downloads


@pytest.fixture
def seeded(ws, member, lead):
    for n, v in [("A", 1), ("B", 2), ("C", 3)]:
        RecordService.create_record(workspace_id=ws.id, member=member, entity=lead,
                                    data={"name": n, "value": v})
    return lead


@pytest.mark.django_db
class TestExport:
    def test_csv_export(self, ws, member, seeded):
        job = ExportService.create_job(entity_slug="lead", format="csv",
                                       requested_by=member.user_id, workspace_id=ws.id)
        job.refresh_from_db()
        assert job.status == "completed"
        assert job.row_count == 3
        assert job.output_storage_key
        assert job.download_expires_at is not None

    def test_xlsx_export_roundtrip(self, ws, member, seeded):
        from openpyxl import load_workbook

        from apps.documents.storage import get_storage_backend
        job = ExportService.create_job(entity_slug="lead", format="xlsx",
                                       requested_by=member.user_id, workspace_id=ws.id)
        job.refresh_from_db()
        data = get_storage_backend(ws.id).download(job.output_storage_key)
        wb = load_workbook(io.BytesIO(data))
        assert wb.active.max_row == 4   # header + 3 rows

    def test_download_url(self, ws, member, seeded):
        job = ExportService.create_job(entity_slug="lead", format="csv",
                                       requested_by=member.user_id, workspace_id=ws.id)
        url = ExportService.download_url(job.id, ws.id)
        assert "sig=" in url

    def test_expire_deletes_artifact(self, ws, member, seeded):
        from datetime import timedelta

        from django.utils import timezone
        job = ExportService.create_job(entity_slug="lead", format="csv",
                                       requested_by=member.user_id, workspace_id=ws.id)
        ExportJob.objects.filter(id=job.id).update(
            download_expires_at=timezone.now() - timedelta(hours=1))
        assert expire_export_downloads() == 1
        job.refresh_from_db()
        assert job.output_storage_key == ""
        with pytest.raises(ImportError):
            ExportService.download_url(job.id, ws.id)
