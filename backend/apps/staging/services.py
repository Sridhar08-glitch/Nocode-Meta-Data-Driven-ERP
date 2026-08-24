"""
Import / Export pipeline (PROJECT_HANDBOOK.md §29.1).

Import: upload → parse (CSV/XLSX) → column mapping → validate (per `FieldMap`)
→ confirm → batched import through `RecordService` (honouring duplicate strategy).
Export: NQL → CSV/XLSX/JSON → storage artifact + signed, expiring download.

File bytes live in the document storage backend (never the DB); only the
``storage_key`` is persisted.
"""
from __future__ import annotations

import csv
import io
import json
import uuid

from django.utils import timezone

from apps.documents.storage import get_storage_backend
from apps.nql.exceptions import NQLError
from apps.nql.services import execute_nql
from apps.records import dal
from apps.records.services import RecordService, resolve_entity
from apps.schema_registry.exceptions import EntityNotFoundError
from apps.schema_registry.services import SchemaRegistryService

from .models import ExportJob, ImportJob, ImportRow

BATCH_SIZE = 100
PREVIEW_ROWS = 50


class ImportError(Exception):  # noqa: A001,N818 — domain error (shadows builtin intentionally)
    pass


def _system_member(workspace_id, user_id):
    from apps.workflows.executors import SystemMember
    actor = user_id or uuid.UUID(int=0)
    return SystemMember(user_id=actor, id=actor)


def _entity(workspace_id, slug):
    try:
        return SchemaRegistryService.get_entity(workspace_id=workspace_id, slug=slug)
    except EntityNotFoundError as exc:
        raise ImportError(str(exc)) from exc


# ── file parsing ──────────────────────────────────────────────────────────────
def _parse_bytes(data: bytes, filename: str, delimiter: str = ",") -> tuple[list[str], list[list]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [[("" if c is None else c) for c in row]
                for row in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        return [], []
    headers = [str(h).strip() for h in rows[0]]
    return headers, rows[1:]


class ImportService:
    @staticmethod
    def create_job(*, entity_slug, file_obj, filename, duplicate_strategy="skip",
                   initiated_by, workspace_id, column_mapping=None,
                   match_field_slug="", delimiter=",") -> ImportJob:
        entity = _entity(workspace_id, entity_slug)
        data = file_obj.read() if hasattr(file_obj, "read") else bytes(file_obj)
        backend = get_storage_backend(workspace_id)
        ws_hex8 = uuid.UUID(str(workspace_id)).hex[:8]
        key = f"{ws_hex8}/imports/{uuid.uuid4().hex}/{filename}"
        backend.upload(data, key)
        job = ImportJob.objects.create(
            workspace_id=workspace_id, entity_id=entity.id, entity_slug=entity_slug,
            initiated_by=initiated_by, status="parsing", source_filename=filename,
            source_storage_key=key, duplicate_strategy=duplicate_strategy,
            match_field_slug=match_field_slug, delimiter=delimiter,
            column_mapping=column_mapping or {})
        from .tasks import parse_import_file
        parse_import_file.delay(str(job.id))
        job.refresh_from_db()
        return job

    @staticmethod
    def parse_file(job_id) -> None:
        job = ImportJob.objects.get(id=job_id)
        backend = get_storage_backend(job.workspace_id)
        data = backend.download(job.source_storage_key)
        headers, rows = _parse_bytes(data, job.source_filename, job.delimiter)
        entity = _entity(job.workspace_id, job.entity_slug)
        field_slugs = set(dal.FieldMap(entity).fields)

        auto = {h: (h if h in field_slugs else None) for h in headers}
        mapping = {**auto, **(job.column_mapping or {})}

        ImportRow.objects.filter(job_id=job.id).delete()
        bulk = []
        for i, raw in enumerate(rows, start=1):
            raw_dict = {headers[j]: (raw[j] if j < len(raw) else "")
                        for j in range(len(headers))}
            bulk.append(ImportRow(job_id=job.id, workspace_id=job.workspace_id,
                                  row_number=i, raw_data=raw_dict, status="pending"))
        ImportRow.objects.bulk_create(bulk, batch_size=500)

        job.column_mapping = mapping
        job.total_rows = len(rows)
        job.status = "validating"
        job.save(update_fields=["column_mapping", "total_rows", "status"])
        ImportService.validate_rows(job.id)

    @staticmethod
    def set_column_mapping(job_id, mapping: dict) -> ImportJob:
        job = ImportJob.objects.get(id=job_id)
        entity = _entity(job.workspace_id, job.entity_slug)
        field_slugs = set(dal.FieldMap(entity).fields)
        for _src, slug in (mapping or {}).items():
            if slug and slug not in field_slugs:
                raise ImportError(f"Unknown field {slug!r}")
        job.column_mapping = {**(job.column_mapping or {}), **mapping}
        job.status = "validating"
        job.save(update_fields=["column_mapping", "status"])
        from .tasks import validate_import
        validate_import.delay(str(job.id))
        job.refresh_from_db()
        return job

    @staticmethod
    def validate_rows(job_id) -> None:
        job = ImportJob.objects.get(id=job_id)
        entity = _entity(job.workspace_id, job.entity_slug)
        fmap = dal.FieldMap(entity)
        mapping = {src: slug for src, slug in (job.column_mapping or {}).items() if slug}
        valid = invalid = 0
        for row in ImportRow.objects.filter(job_id=job.id).iterator(chunk_size=500):
            mapped, errors = {}, []
            for src, slug in mapping.items():
                fd = fmap.fields.get(slug)
                if fd is None:
                    continue
                raw = (row.raw_data or {}).get(src)
                if fd.is_required and (raw is None or str(raw).strip() == ""):
                    errors.append({"field": slug, "message": "required"})
                    continue
                if raw in (None, ""):
                    continue
                try:
                    mapped[slug] = fmap.coerce(slug, raw)
                except NQLError:
                    errors.append({"field": slug, "message": "invalid value"})
            # required fields not present in the mapping at all
            for slug, fd in fmap.fields.items():
                if fd.is_required and not fd.is_system and slug not in mapped \
                        and not any(e["field"] == slug for e in errors):
                    errors.append({"field": slug, "message": "required"})
            row.mapped_data = mapped
            row.validation_errors = errors
            row.status = "invalid" if errors else "valid"
            row.save(update_fields=["mapped_data", "validation_errors", "status"])
            if errors:
                invalid += 1
            else:
                valid += 1
        job.valid_rows = valid
        job.invalid_rows = invalid
        job.status = "awaiting_confirm"
        job.save(update_fields=["valid_rows", "invalid_rows", "status"])

    @staticmethod
    def confirm_import(job_id, actor_id=None) -> ImportJob:
        job = ImportJob.objects.get(id=job_id)
        if job.status != "awaiting_confirm":
            raise ImportError(f"Job is not awaiting confirmation (status={job.status})")
        job.status = "importing"
        job.started_at = timezone.now()
        job.save(update_fields=["status", "started_at"])
        from .tasks import execute_import
        execute_import.delay(str(job.id))
        job.refresh_from_db()
        return job

    @staticmethod
    def execute_import(job_id) -> None:
        job = ImportJob.objects.get(id=job_id)
        entity = resolve_entity(job.workspace_id, job.entity_slug)
        member = _system_member(job.workspace_id, job.initiated_by)
        imported = skipped = errored = 0
        rows = ImportRow.objects.filter(job_id=job.id, status="valid").order_by("row_number")
        for row in rows.iterator(chunk_size=BATCH_SIZE):
            try:
                existing = ImportService._find_duplicate(job, entity, member, row.mapped_data)
                if existing is not None:
                    if job.duplicate_strategy == "skip":
                        row.status = "skipped"
                        row.save(update_fields=["status"])
                        skipped += 1
                        continue
                    if job.duplicate_strategy == "error":
                        row.status = "error"
                        row.validation_errors = [{"field": job.match_field_slug,
                                                  "message": "duplicate"}]
                        row.save(update_fields=["status", "validation_errors"])
                        errored += 1
                        continue
                    RecordService.update_record(
                        workspace_id=job.workspace_id, member=member, entity=entity,
                        record_id=existing, data=row.mapped_data)
                    rid = existing
                else:
                    created = RecordService.create_record(
                        workspace_id=job.workspace_id, member=member, entity=entity,
                        data=row.mapped_data)
                    rid = created.get("id")
                row.status = "imported"
                row.imported_record_id = rid
                row.save(update_fields=["status", "imported_record_id"])
                imported += 1
            except Exception as exc:  # noqa: BLE001 — one bad row must not abort the batch
                row.status = "error"
                row.validation_errors = [{"field": "", "message": str(exc)}]
                row.save(update_fields=["status", "validation_errors"])
                errored += 1
            ImportJob.objects.filter(id=job.id).update(imported_rows=imported,
                                                       skipped_rows=skipped,
                                                       error_rows=errored)
        job.status = "completed"
        job.completed_at = timezone.now()
        job.imported_rows = imported
        job.skipped_rows = skipped
        job.error_rows = errored
        job.save(update_fields=["status", "completed_at", "imported_rows",
                                "skipped_rows", "error_rows"])

    @staticmethod
    def _find_duplicate(job, entity, member, mapped_data):
        if not job.match_field_slug:
            return None
        value = mapped_data.get(job.match_field_slug)
        if value in (None, ""):
            return None
        query = {"entity": entity.slug, "filter": {"op": "and", "conditions": [
            {"field": job.match_field_slug, "op": "=", "value": str(value)}]}, "limit": 1}
        try:
            rows = execute_nql(workspace_id=job.workspace_id, source=query,
                               user_id=member.user_id)
        except NQLError:
            return None
        return rows[0]["id"] if rows else None

    @staticmethod
    def preview(job_id, limit=PREVIEW_ROWS):
        return ImportRow.objects.filter(job_id=job_id).order_by("row_number")[:limit]

    @staticmethod
    def cancel(job_id) -> ImportJob:
        job = ImportJob.objects.get(id=job_id)
        if job.status not in ("completed", "failed"):
            job.status = "cancelled"
            job.save(update_fields=["status"])
        return job


class ExportService:
    @staticmethod
    def create_job(*, entity_slug=None, nql_ast=None, format="csv", requested_by,
                   workspace_id, include_fields=None, report_id=None) -> ExportJob:
        entity = _entity(workspace_id, entity_slug) if entity_slug else None
        if nql_ast is None and entity is not None:
            nql_ast = {"entity": entity_slug}
        job = ExportJob.objects.create(
            workspace_id=workspace_id, entity_id=entity.id if entity else None,
            report_id=report_id, initiated_by=requested_by, status="queued",
            nql_ast=nql_ast, format=format, include_fields=include_fields or [])
        from .tasks import execute_export
        execute_export.delay(str(job.id))
        job.refresh_from_db()
        return job

    @staticmethod
    def execute_export(job_id) -> None:
        import datetime as _dt
        job = ExportJob.objects.get(id=job_id)
        job.status = "running"
        job.started_at = timezone.now()
        job.save(update_fields=["status", "started_at"])
        try:
            rows = execute_nql(workspace_id=job.workspace_id, source=job.nql_ast,
                               user_id=job.initiated_by)
            if job.include_fields:
                rows = [{k: r.get(k) for k in job.include_fields} for r in rows]
            data = ExportService._serialize(rows, job.format)
            backend = get_storage_backend(job.workspace_id)
            ws_hex8 = uuid.UUID(str(job.workspace_id)).hex[:8]
            filename = f"export-{job.id}.{job.format}"
            key = f"{ws_hex8}/exports/{uuid.uuid4().hex}/{filename}"
            backend.upload(data, key)
            job.output_storage_key = key
            job.output_filename = filename
            job.row_count = len(rows)
            job.size_bytes = len(data)
            job.download_expires_at = timezone.now() + _dt.timedelta(hours=24)
            job.status = "completed"
            job.completed_at = timezone.now()
            job.save(update_fields=["output_storage_key", "output_filename", "row_count",
                                    "size_bytes", "download_expires_at", "status",
                                    "completed_at"])
            ExportService._notify(job)
        except Exception as exc:  # noqa: BLE001
            job.status = "failed"
            job.error_message = str(exc)
            job.save(update_fields=["status", "error_message"])

    @staticmethod
    def _serialize(rows, fmt) -> bytes:
        keys = list(rows[0].keys()) if rows else []
        if fmt == "json":
            return json.dumps(rows, default=str).encode("utf-8")
        if fmt == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(keys)
            for r in rows:
                ws.append([r.get(k) for k in keys])
            out = io.BytesIO()
            wb.save(out)
            return out.getvalue()
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(keys)
        for r in rows:
            writer.writerow([r.get(k, "") for k in keys])
        return buf.getvalue().encode("utf-8")

    @staticmethod
    def _notify(job):
        try:
            from apps.notifications.services import NotificationService
            NotificationService.send(
                recipient_id=job.initiated_by, recipient_type="member",
                template_slug="export_ready",
                context={"event_type": "export_ready", "subject": "Export ready",
                         "body": f"Your export ({job.row_count} rows) is ready to download."},
                workspace_id=job.workspace_id, channels=["in_app"])
        except Exception:  # noqa: BLE001
            pass

    @staticmethod
    def download_url(job_id, workspace_id):
        job = ExportJob.objects.filter(id=job_id, workspace_id=workspace_id).first()
        if job is None or job.status != "completed" or not job.output_storage_key:
            raise ImportError("Export not available")
        if job.download_expires_at and job.download_expires_at < timezone.now():
            raise ImportError("Download link expired")
        backend = get_storage_backend(workspace_id)
        return backend.generate_signed_url(job.output_storage_key, 300)
