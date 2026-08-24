"""ReportService (PROJECT_HANDBOOK.md §25.1 / §25.4)."""
import io

import pytest

from apps.reporting import services as svc
from apps.reporting.models import ReportSnapshot
from apps.reporting.services import ReportError, ReportService


@pytest.mark.django_db
class TestExecute:
    def test_execute_table(self, ws, leads, make_report):
        report = make_report()
        result = ReportService.execute_report(report, ws.id)
        assert result["total_count"] == 3
        assert result["truncated"] is False
        assert any(c["key"] == "status" for c in result["columns"])

    def test_column_labels_and_order(self, ws, leads, make_report):
        report = make_report(display_config={
            "columns": [{"key": "name", "label": "Lead Name"}, {"key": "status"}]})
        result = ReportService.execute_report(report, ws.id)
        assert [c["key"] for c in result["columns"]] == ["name", "status"]
        assert result["columns"][0]["label"] == "Lead Name"

    def test_truncation(self, ws, leads, make_report, monkeypatch):
        monkeypatch.setattr(svc, "MAX_ROWS", 2)
        report = make_report()
        result = ReportService.execute_report(report, ws.id)
        assert result["truncated"] is True
        assert result["total_count"] == 2


@pytest.mark.django_db
class TestPivot:
    def test_sum_by_status(self, ws, leads, make_report):
        report = make_report(report_type="pivot", display_config={
            "row_field": "status", "value_field": "value", "agg": "sum"})
        result = ReportService.execute_pivot(report, ws.id)
        totals = dict(zip(result["row_values"],
                          [row[0] for row in result["matrix"]], strict=False))
        assert totals["open"] == 300.0
        assert totals["won"] == 50.0

    def test_invalid_pivot_config(self, ws, leads, make_report):
        report = make_report(report_type="pivot", display_config={})
        with pytest.raises(ReportError):
            ReportService.execute_pivot(report, ws.id)


@pytest.mark.django_db
class TestSnapshot:
    def test_snapshot_stores_rows(self, ws, leads, make_report):
        report = make_report()
        snap = ReportService.take_snapshot(report, ws.id)
        assert snap.row_count == 3
        assert len(snap.data) == 3

    def test_snapshot_pruned_to_five(self, ws, leads, make_report):
        report = make_report()
        for _ in range(7):
            ReportService.take_snapshot(report, ws.id)
        assert ReportSnapshot.objects.filter(report_id=report.id).count() == 5


@pytest.mark.django_db
class TestExport:
    def test_csv(self, ws, leads, make_report):
        report = make_report(display_config={
            "columns": [{"key": "name", "label": "Name"}, {"key": "value", "label": "Value"}]})
        data = ReportService.export_to_csv(report, ws.id).decode()
        lines = [ln for ln in data.splitlines() if ln]
        assert lines[0] == "Name,Value"
        assert len(lines) == 4   # header + 3 rows

    def test_xlsx(self, ws, leads, make_report):
        from openpyxl import load_workbook
        report = make_report(display_config={
            "columns": [{"key": "name", "label": "Name"}, {"key": "value", "label": "Value"}]})
        data = ReportService.export_to_xlsx(report, ws.id)
        wb = load_workbook(io.BytesIO(data))
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[0] == ("Name", "Value")
        assert len(rows) == 4


class TestValidateNql:
    def test_valid_source(self):
        assert ReportService.validate_nql(nql_source="FROM lead WHERE value > 100") == []

    def test_invalid_source(self):
        errors = ReportService.validate_nql(nql_source="FROM lead WHERE ((")
        assert errors

    def test_valid_ast(self):
        assert ReportService.validate_nql(nql_ast={"entity": "lead"}) == []

    def test_missing_input(self):
        assert ReportService.validate_nql() == ["Provide nql_source or nql_ast"]
